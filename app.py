"""
FoldNote — anteckningar där varje stycke kan ha flera alternativa formuleringar.

All data ligger i en enda JSON-fil: data/foldnote.json. Den är läsbar,
går att versionshantera, redigera för hand och kopiera var som helst.

Skrivningar är atomära: filen skrivs först till en temporär fil i samma
mapp och byts sedan in med os.replace, som är en atomär operation på
Linux, macOS och Windows. Ett strömavbrott mitt i en skrivning kan
alltså inte lämna en halv fil efter sig — antingen finns den gamla
versionen kvar, eller den nya.
"""

import json
import os
import shutil
import threading
import uuid
from datetime import datetime
from pathlib import Path

from flask import Flask, jsonify, render_template, request, send_file, send_from_directory

BASE = Path(__file__).resolve().parent
DATA = BASE / "data"
STORE = DATA / "foldnote.json"
LEGACY_XLSX = DATA / "foldnote.xlsx"
LEGACY_DB = DATA / "foldnote.db"

app = Flask(__name__)

_lock = threading.Lock()


# --------------------------------------------------------------------------
# Lagring
# --------------------------------------------------------------------------

def iso():
    return datetime.now().isoformat(timespec="seconds")


def new_id():
    return str(uuid.uuid4())


def empty_store():
    return {"version": 1, "notes": []}


def load_store():
    """Läser hela filen. En trasig fil läggs undan i stället för att raderas."""
    if not STORE.exists():
        return empty_store()
    try:
        with STORE.open("r", encoding="utf-8") as f:
            data = json.load(f)
    except (json.JSONDecodeError, OSError) as exc:
        broken = DATA / f"foldnote-trasig-{datetime.now():%Y%m%d-%H%M%S}.json"
        shutil.copy2(STORE, broken)
        print(f"Kunde inte läsa {STORE.name} ({exc}). Kopia sparad som {broken.name}.")
        return empty_store()

    if not isinstance(data, dict) or not isinstance(data.get("notes"), list):
        return empty_store()
    return data


def save_store(data):
    """Skriver hela filen atomärt."""
    DATA.mkdir(exist_ok=True)
    tmp = STORE.with_suffix(".json.tmp")
    with tmp.open("w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
        f.flush()
        os.fsync(f.fileno())
    os.replace(tmp, STORE)


def find_note(data, nid):
    for note in data["notes"]:
        if note.get("id") == nid:
            return note
    return None


def normalize(note):
    """Ser till att en anteckning har alla fält, oavsett var den kom ifrån."""
    note.setdefault("id", new_id())
    note.setdefault("title", "")
    note.setdefault("tags", "")
    note.setdefault("created_at", iso())
    note.setdefault("updated_at", note["created_at"])
    segments = note.get("segments")
    note["segments"] = segments if isinstance(segments, list) else []
    for seg in note["segments"]:
        seg.setdefault("id", new_id())
        seg.setdefault("original", "")
        versions = seg.get("versions")
        seg["versions"] = versions if isinstance(versions, list) else []
        for ver in seg["versions"]:
            ver.setdefault("id", new_id())
            ver.setdefault("label", "")
            ver.setdefault("text", "")
    return note


# --------------------------------------------------------------------------
# Migrering från tidigare format
# --------------------------------------------------------------------------

def migrate():
    """Flyttar över data från en tidigare foldnote.db eller foldnote.xlsx."""
    DATA.mkdir(exist_ok=True)
    if STORE.exists():
        return

    notes = _from_sqlite() or _from_xlsx()
    if not notes:
        return

    save_store({"version": 1, "notes": [normalize(n) for n in notes]})
    print(f"Migrerade {len(notes)} anteckningar till {STORE.name}.")


def _from_sqlite():
    if not LEGACY_DB.exists():
        return None
    import sqlite3

    try:
        conn = sqlite3.connect(LEGACY_DB)
        conn.row_factory = sqlite3.Row
        notes = []
        for n in conn.execute("SELECT * FROM notes ORDER BY updated_at DESC").fetchall():
            segments = []
            for s in conn.execute(
                "SELECT * FROM segments WHERE note_id = ? ORDER BY pos", (n["id"],)
            ).fetchall():
                versions = conn.execute(
                    "SELECT id, label, text FROM versions WHERE segment_id = ? ORDER BY pos",
                    (s["id"],),
                ).fetchall()
                segments.append(
                    {
                        "id": s["id"],
                        "original": s["original"],
                        "versions": [dict(v) for v in versions],
                    }
                )
            notes.append(
                {
                    "id": n["id"],
                    "title": n["title"],
                    "tags": n["tags"],
                    "created_at": n["created_at"],
                    "updated_at": n["updated_at"],
                    "segments": segments,
                }
            )
        conn.close()
        LEGACY_DB.rename(DATA / "foldnote-gammal.db")
        return notes
    except Exception as exc:
        print(f"Kunde inte läsa foldnote.db ({exc}). Filen lämnas orörd.")
        return None


def _from_xlsx():
    if not LEGACY_XLSX.exists():
        return None
    try:
        from openpyxl import load_workbook
    except ImportError:
        print(
            "data/foldnote.xlsx finns men openpyxl är inte installerat. "
            "Kör `pip install openpyxl` en gång och starta om, så flyttas datan över."
        )
        return None

    try:
        book = load_workbook(LEGACY_XLSX, read_only=True)

        def rows(name):
            return list(book[name].iter_rows(min_row=2, values_only=True))

        versions_by_segment = {}
        for r in rows("Versions"):
            if r and r[0]:
                versions_by_segment.setdefault(r[1], []).append(
                    {"id": r[0], "pos": r[2] or 0, "label": r[4] or "", "text": r[3] or ""}
                )

        segments_by_note = {}
        for r in rows("Segments"):
            if r and r[0]:
                vers = sorted(versions_by_segment.get(r[0], []), key=lambda v: v["pos"])
                for v in vers:
                    v.pop("pos", None)
                segments_by_note.setdefault(r[1], []).append(
                    {"id": r[0], "pos": r[2] or 0, "original": r[3] or "", "versions": vers}
                )

        notes = []
        for r in rows("Notes"):
            if not r or not r[0]:
                continue
            segs = sorted(segments_by_note.get(r[0], []), key=lambda s: s["pos"])
            for s in segs:
                s.pop("pos", None)
            notes.append(
                {
                    "id": r[0],
                    "title": r[1] or "",
                    "tags": r[4] or "",
                    "created_at": r[2] or iso(),
                    "updated_at": r[3] or iso(),
                    "segments": segs,
                }
            )
        book.close()
        LEGACY_XLSX.rename(DATA / "foldnote-gammal.xlsx")
        return notes
    except Exception as exc:
        print(f"Kunde inte läsa foldnote.xlsx ({exc}). Filen lämnas orörd.")
        return None


# --------------------------------------------------------------------------
# Sidor
# --------------------------------------------------------------------------

@app.get("/")
def index():
    return render_template("index.html")


@app.get("/manifest.json")
def manifest():
    return send_from_directory("static", "manifest.json", mimetype="application/manifest+json")


@app.get("/sw.js")
def service_worker():
    # Måste ligga i roten — från /static/ får service workern bara
    # scope över /static/ och kan inte kontrollera appen.
    return send_from_directory("static", "sw.js", mimetype="application/javascript")


# --------------------------------------------------------------------------
# API
# --------------------------------------------------------------------------

@app.get("/api/notes")
def list_notes():
    data = load_store()
    out = []
    for note in data["notes"]:
        segments = note.get("segments") or []
        first = segments[0].get("original", "") if segments else ""
        out.append(
            {
                "id": note.get("id"),
                "title": note.get("title", ""),
                "tags": note.get("tags", ""),
                "updated_at": note.get("updated_at", ""),
                "preview": (first or "")[:120],
            }
        )
    out.sort(key=lambda n: n["updated_at"] or "", reverse=True)
    return jsonify(out)


@app.post("/api/notes")
def create_note():
    payload = request.get_json(silent=True) or {}
    now = iso()
    note = normalize(
        {
            "id": new_id(),
            "title": payload.get("title", ""),
            "tags": payload.get("tags", ""),
            "created_at": now,
            "updated_at": now,
            "segments": [{"id": new_id(), "original": "", "versions": []}],
        }
    )
    with _lock:
        data = load_store()
        data["notes"].insert(0, note)
        save_store(data)
    return jsonify(note), 201


@app.get("/api/notes/<nid>")
def get_note(nid):
    note = find_note(load_store(), nid)
    if note is None:
        return jsonify({"error": "Anteckningen finns inte."}), 404
    return jsonify(normalize(note))


@app.put("/api/notes/<nid>")
def save_note(nid):
    payload = request.get_json(silent=True) or {}
    with _lock:
        data = load_store()
        note = find_note(data, nid)
        if note is None:
            return jsonify({"error": "Anteckningen finns inte."}), 404

        note["title"] = payload.get("title", "")
        note["tags"] = payload.get("tags", "")
        note["updated_at"] = iso()
        note["segments"] = [
            {
                "id": seg.get("id") or new_id(),
                "original": seg.get("original", ""),
                "versions": [
                    {
                        "id": ver.get("id") or new_id(),
                        "label": ver.get("label", ""),
                        "text": ver.get("text", ""),
                    }
                    for ver in (seg.get("versions") or [])
                ],
            }
            for seg in (payload.get("segments") or [])
        ]
        save_store(data)
    # Tillbaka med hela anteckningen så klienten får riktiga id:n på
    # nya stycken och versioner.
    return jsonify(note)


@app.delete("/api/notes/<nid>")
def delete_note(nid):
    with _lock:
        data = load_store()
        before = len(data["notes"])
        data["notes"] = [n for n in data["notes"] if n.get("id") != nid]
        if len(data["notes"]) == before:
            return jsonify({"error": "Anteckningen finns inte."}), 404
        save_store(data)
    return jsonify({"ok": True})


@app.get("/api/export")
def export():
    """Laddar ner hela datafilen som den ser ut."""
    if not STORE.exists():
        save_store(empty_store())
    return send_file(
        STORE,
        as_attachment=True,
        download_name=f"foldnote-{datetime.now():%Y-%m-%d}.json",
        mimetype="application/json",
    )


@app.post("/api/import")
def import_notes():
    """Lägger till anteckningar från en tidigare export. Skriver inte över
    något befintligt — allt som kommer in får nya id:n."""
    payload = request.get_json(silent=True)
    incoming = payload.get("notes") if isinstance(payload, dict) else payload
    if not isinstance(incoming, list):
        return jsonify({"error": "Filen innehåller inga anteckningar."}), 400

    added = []
    for raw in incoming:
        if not isinstance(raw, dict):
            continue
        note = normalize(dict(raw))
        note["id"] = new_id()
        for seg in note["segments"]:
            seg["id"] = new_id()
            for ver in seg["versions"]:
                ver["id"] = new_id()
        added.append(note)

    if not added:
        return jsonify({"error": "Filen innehåller inga anteckningar."}), 400

    with _lock:
        data = load_store()
        data["notes"] = added + data["notes"]
        save_store(data)
    return jsonify({"ok": True, "added": len(added)})


@app.errorhandler(404)
def not_found(_e):
    if request.path.startswith("/api/"):
        return jsonify({"error": "Hittades inte."}), 404
    return "Sidan hittades inte.", 404


migrate()

if __name__ == "__main__":
    # Debugläget exponerar Werkzeugs felsökare på hela nätverket eftersom
    # servern lyssnar på 0.0.0.0. Sätt FLASK_DEBUG=0 för att stänga av.
    app.run(
        host="0.0.0.0",
        port=int(os.environ.get("PORT", 5000)),
        debug=os.environ.get("FLASK_DEBUG", "1") == "1",
    )
